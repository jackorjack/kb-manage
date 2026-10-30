"""In-process document to Markdown conversion (docx / xlsx / csv / md)."""

from __future__ import annotations

import csv
from pathlib import Path
from typing import List, Optional

from markitdown import MarkItDown

from .files import FileValidationError, normalize_markdown


def _csv_to_markdown(data: bytes) -> str:
    """CSV -> per-record Markdown lines that chunk well for retrieval."""
    text = data.decode("utf-8-sig", errors="replace")
    rows = [row for row in csv.reader(text.splitlines()) if any(c.strip() for c in row)]
    if not rows:
        return ""
    header, body = rows[0], rows[1:]
    lines = [f"# {len(body)} 条记录", ""]
    for row in body:
        parts = [f"{h.strip()}: {c.strip()}" for h, c in zip(header, row) if c.strip()]
        if parts:
            lines.append("- " + " | ".join(parts))
    return "\n".join(lines)


def _cell(value) -> str:
    return "" if value is None else str(value).strip()


def _columnar_sheet(rows: List[List]) -> bool:
    """Detect a metadata-headed columnar sheet (like the keyword matrix).

    Shape: a sparse label column (A) that only fills a few leading metadata
    rows, while the remaining columns carry long independent lists. A plain
    2-D table keeps column A populated throughout, so this stays False for it.
    """
    if len(rows) < 4 or not rows or len(rows[0]) < 2:
        return False
    a_filled_rows = [i for i, row in enumerate(rows) if row and _cell(row[0])]
    if not a_filled_rows:
        return False
    # Column A values cluster at the top (metadata), sparse everywhere else.
    total = len(rows)
    a_filled = len(a_filled_rows)
    top_filled = sum(1 for i in a_filled_rows if i < 10)
    other_filled = sum(1 for row in rows if any(_cell(c) for c in row[1:]))
    return (
        top_filled >= 2
        and a_filled <= total * 0.25
        and other_filled >= total * 0.3
    )


def _columnar_structure(rows: List[List]) -> tuple[List[List], int]:
    """Split a columnar sheet into (metadata_rows, header_index).

    The header row is the first row with an empty label column that still
    carries column names. Metadata rows are every row that fills the label
    column (A) — they describe the columns rather than holding a record.
    """
    header_index = next(
        (i for i, row in enumerate(rows) if row and not _cell(row[0]) and any(_cell(c) for c in row[1:])),
        0,
    )
    meta_rows = [row for row in rows if row and _cell(row[0])]
    return meta_rows, header_index


def _columnar_to_markdown(sheet_title: str, rows: List[List]) -> str:
    """Matrix-style sheet: metadata head rows, then each column listed on its
    own so unrelated same-row values are never glued into a fake record."""
    meta_rows, header_index = _columnar_structure(rows)
    # Body = rows after the last metadata row; a label-column row like
    # "GEO问题/标题" still counts as metadata, so data starts right after it.
    last_meta = max(
        (i for i, row in enumerate(rows) if row and _cell(row[0])),
        default=header_index,
    )
    body_rows = rows[last_meta + 1 :]
    header_row = rows[header_index]
    col_names = [_cell(c) or f"列{i+1}" for i, c in enumerate(header_row)]

    lines = [f"## {sheet_title}", ""]
    # Emit metadata rows (label column A + its values) as readable context.
    for row in meta_rows:
        label = _cell(row[0]) if row else ""
        values = " | ".join(_cell(c) for c in row[1:] if _cell(c))
        if label and values:
            lines.append(f"**{label}**: {values}")
        elif label:
            lines.append(f"**{label}**")
    if meta_rows:
        lines.append("")

    # Then each column becomes its own bullet list of independent values.
    body = body_rows
    for col_index, name in enumerate(col_names):
        if col_index == 0:
            continue
        values = [
            _cell(row[col_index])
            for row in body
            if col_index < len(row) and _cell(row[col_index])
        ]
        if not values:
            continue
        lines.append(f"### {name}")
        for value in values:
            lines.append(f"- {value}")
        lines.append("")
    return "\n".join(lines)


def _table_to_markdown(sheet_title: str, rows: List[List]) -> str:
    """Plain 2-D sheet: first row is the header, each later row one record."""
    header = [_cell(c) or f"列{i+1}" for i, c in enumerate(rows[0])]
    lines = [f"## {sheet_title}", ""]
    for row in rows[1:]:
        parts = [f"- **{h}**: {c}" for h, c in zip(header, row) if _cell(c)]
        if parts:
            lines.append("\n".join(parts) + "\n")
    return "\n".join(lines)


def _xlsx_to_markdown(source: Path) -> str:
    """XLSX -> structured Markdown.

    Two layouts are handled per sheet:
    - matrix/columnar sheets (metadata head + independent columns) are
      flattened column-by-column, which suits keyword/taxonomy tables;
    - plain 2-D tables keep the row-per-record `field: value` shape.
    """
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise FileValidationError("xlsx support requires the openpyxl package") from exc
    try:
        workbook = load_workbook(source, read_only=True, data_only=True)
    except Exception as exc:
        raise FileValidationError(f"无法读取 Excel 文件: {exc}") from exc
    sections = []
    for sheet in workbook.worksheets:
        rows = [list(r) for r in sheet.iter_rows(values_only=True)]
        rows = [r for r in rows if any(_cell(c) for c in r)]
        if not rows:
            continue
        if _columnar_sheet(rows):
            sections.append(_columnar_to_markdown(sheet.title, rows))
        else:
            sections.append(_table_to_markdown(sheet.title, rows))
    if not sections:
        raise FileValidationError("Excel 文件没有可读内容")
    return "\n".join(sections)


class DocxMarkdownConverter:
    def __init__(self, max_output_bytes: int):
        self.max_output_bytes = max_output_bytes
        self.converter = MarkItDown(enable_plugins=False)

    def convert(self, source: Path, output: Path, extension: str) -> None:
        if extension == ".md":
            data = source.read_bytes()
            try:
                markdown = data.decode("utf-8")
            except UnicodeDecodeError as exc:
                raise FileValidationError("Markdown file must be UTF-8") from exc
        elif extension == ".docx":
            with source.open("rb") as stream:
                result = self.converter.convert_stream(stream, file_extension=".docx")
            markdown = result.markdown
        elif extension == ".xlsx":
            markdown = _xlsx_to_markdown(source)
        elif extension == ".csv":
            markdown = _csv_to_markdown(source.read_bytes())
        else:
            raise FileValidationError("unsupported source extension")
        output.parent.mkdir(parents=True, exist_ok=True)
        output.write_bytes(normalize_markdown(markdown, self.max_output_bytes))
